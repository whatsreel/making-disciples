#!/usr/bin/env python3
"""Drift check: does every view still agree with master.json?

Run this before circulating anything. It is the guard against the failure
this project has already had twice — a view edited by hand and then silently
disagreeing with the others.
"""
import os
WORK = os.environ.get('WORK', '/home/claude')
OUT = os.environ.get('OUT', '/mnt/user-data/outputs')
import json, re, sys

M = json.load(open(WORK+'/master.json', encoding='utf-8'))
args = M['arguments']
ids = {a['id'] for a in args}
problems = []

def check(label, ok, detail=''):
    print(f'  {"PASS" if ok else "FAIL"}  {label}{"" if ok else "  -> " + detail}')
    if not ok:
        problems.append(label)

print('master.json')
check(f'{len(args)} arguments, unique ids', len(ids) == len(args))
check('every argument has verbatim prose', all(a.get('body_md') for a in args))
check('every argument has a verdict', all(a.get('lands_verdict') for a in args))

print('\noutputs/master.json  (the copy the person sees)')
try:
    import filecmp
    same = filecmp.cmp(WORK+'/master.json', OUT+'/master.json', shallow=False)
    check('outputs copy identical to working master', same, 'stale copy — re-run cp')
except FileNotFoundError:
    check('outputs copy present', False, 'missing')

print('\nargument-index.md  (prose source)')
idx = open(OUT+'/argument-index.md', encoding='utf-8').read()
notv = [a['id'] for a in args if a['body_md'] not in idx]
check('all prose verbatim in index', not notv, str(notv))
hdrs = set(re.findall(r'^### ([A-G]\d+[a-z]?)\.', idx, re.M))
check('index headers match master ids', hdrs == ids, str(hdrs ^ ids))

# Each index entry ends with an italic tag line ("*Column: Precept · ... · Who does it: Officers*")
# that readers see on the argument's sheet; args.json holds the same tags as data, and every count
# on the page comes from the data. On 2026-10-04 eleven tags disagreed (C1b-C1f column, C10/C11
# column, what and who), so one sheet said two different things. The two must match field for field.
SURV = {'Holds without': 'Yes', 'Partly holds without': 'Partly', 'Needs Jesus': 'No'}
def _norm(s):
    return (s or '—').replace('’', "'").strip()   # only the apostrophe varies by design
def tag_fields(line):
    """(fields, problems). Every field defaults to absent, so an omitted one is compared too; an
    unrecognized or repeated part is a problem, because readers see it even if the data never does."""
    parts = [p.strip() for p in line.split('·')]
    out = {'band_name': parts[0], 'what': '—', 'who': '—', 'use': '—', 'survives': '—'}
    seen, problems = set(), []
    labels = {'What it is:': 'what', 'Who does it:': 'who', 'Use:': 'use'}
    for p in parts[1:]:
        key = next((v for k, v in labels.items() if p.startswith(k)), None)
        if key:
            val = p.split(':', 1)[1].strip()
        else:
            key = 'survives'
            val = next((v for k, v in sorted(SURV.items(), key=lambda kv: -len(kv[0]))
                        if p.replace('’', "'").startswith(k)), None)
            if val is None:
                problems.append(f"unrecognized tag part '{p}'")
                continue
        if key in seen:
            problems.append(f"repeated {key} ('{p}')")
        seen.add(key)
        out[key] = val
    return out, problems
tag_bad = []
banded = [a for a in args if a.get('band_name') is not None]   # method entries (A) have no column
for a in args:
    if a.get('band_name') is None and a.get('column_line'):
        tag_bad.append(f"{a['id']}: the index gives it a column tag but args.json gives it no column")
    n_tags = len(re.findall(r'^\*Column:', a.get('body_md') or '', re.M))
    if n_tags > 1:   # build_master2 compares only the first; readers see every one
        tag_bad.append(f"{a['id']}: {n_tags} '*Column:' tag lines in one entry")
for a in banded:
    if not a.get('column_line'):
        tag_bad.append(f"{a['id']}: no '*Column: ...*' tag line in the index")
        continue
    t, probs = tag_fields(a['column_line'])
    tag_bad += [f"{a['id']}: {p}" for p in probs]
    for k in ('band_name', 'what', 'who', 'use', 'survives'):
        if _norm(t[k]) != _norm(a.get(k)):
            tag_bad.append(f"{a['id']} {k}: index '{t[k]}' vs args.json '{a.get(k)}'")
print(f'  COUNT tag-lines expected {len(banded)}, disagreeing fields {len(tag_bad)}')
check('index tag lines agree with args.json', not tag_bad, '; '.join(tag_bad[:12]))

print('\nargument-map.html  (board view)')
try:
    html = open(OUT+'/argument-map.html', encoding='utf-8').read()
    hids = set(re.findall(r'id:\s*"([A-G]\d+[a-z]?)"', html)) | set(re.findall(r'"id":\s*"([A-G]\d+[a-z]?)"', html))
    check('map ids match master', hids == ids, str(hids ^ ids))
except FileNotFoundError:
    check('map present', False, 'argument-map.html not found')

try:
    import re as _r
    _m = open(OUT+'/argument-map.html', encoding='utf-8').read()
    _n = len(set(_r.findall(r'"id":\s*"([A-G]\d+[a-z]?)"', _m)) | set(_r.findall(r'\{id:"([A-G]\d+[a-z]?)"', _m)))
    # the template writes the count as 'All '+ARGS.length; a typed 'All 45' in a string or in markup fails
    _hard = _r.search(r"[>'\"]All (\d+)", _m)
    check('map count is derived, not hardcoded', _hard is None, f'hardcoded {_hard.group(1) if _hard else ""}')
except Exception as _e:
    check('map count check ran', False, str(_e))

print('\nargument-map.xlsx  (matrix view)')
try:
    from openpyxl import load_workbook
    ws = load_workbook(OUT+'/argument-map.xlsx')['Positions']
    xids = {r[0] for r in ws.iter_rows(min_row=2, values_only=True) if r[0]}
    check('sheet ids match master', xids == ids, str(xids ^ ids))
except Exception as e:
    check('spreadsheet readable', False, str(e))

print('\ndiscipleship-session.md  (paper view; private, lives in the notes repo)')
# The paper is not published. paper.py finds it in this repo's docs/ or in the notes folder the
# gitignored .notes-expected marker names. The marker says "this machine has the notes": then an
# absent paper FAILS. Without the marker (CI, a public clone) it is counted and passes.
from paper import find_paper, marker_problem, PROBES
paper_path, expected = find_paper(os.environ.get('REPO', os.path.dirname(os.path.abspath(OUT))))
if paper_path:
    paper = open(paper_path, encoding='utf-8').read()
    absent = [k for k, v in PROBES.items() if v not in paper]
    check('recent arguments present in paper', not absent, str(absent))
else:
    print(f'  COUNT paper-present 0/1 ({"marker present" if expected else "no .notes-expected marker: skipped"})')
    if expected:
        check('paper present (this machine has .notes-expected)', False,
              'not in docs/ here, and ' + (marker_problem(os.environ.get('REPO', os.path.dirname(os.path.abspath(OUT))))
                                           or 'the folder the marker names has no docs/discipleship-session.md'))

print('\ntone  (safe to forward to the whole session)')
import tone   # one regex, shared with the gate's check 20
tone_files = [('index', OUT+'/argument-index.md'),
              ('hub template', WORK+'/hub_template.html'),
              ('built page', OUT+'/argument-map.html'),
              ('args.json', WORK+'/args.json'),
              ('front door', WORK+'/front_door.json')]
if paper_path:
    tone_files.insert(1, ('paper', paper_path))
for label, path in tone_files:
    txt = open(path, encoding='utf-8').read()
    if path.endswith('.html') and 'id="payload"' in txt:
        txt = tone.strip_scripture(txt)
    h = tone.hits(txt)
    check(f'{label}: no advocate-vs-adversary framing', not h, str(sorted(set(p.lower() for _, p in h))))

print()
if problems:
    print(f'{len(problems)} problem(s): {problems}')
    sys.exit(1)
print('No drift. All views agree with master.json.')
